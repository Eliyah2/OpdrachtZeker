"""Valideert OpdrachtZeker-financieel-model.xlsx: structuur, kernformules en verwijzingen."""
import re
from openpyxl import load_workbook

wb = load_workbook("OpdrachtZeker-financieel-model.xlsx")
checks = []

def check(name, ok, detail=""):
    checks.append((name, ok, detail))

sheets = wb.sheetnames
check("Bladen aanwezig", sheets == ["Leeswijzer", "Aannames", "Model", "Scenario", "Fasen"], str(sheets))

wa, wm, ws4, wf = wb["Aannames"], wb["Model"], wb["Scenario"], wb["Fasen"]

# Aannames: waarden (rijen volgens generatiescript ORDER)
check("Aannames startkapitaal=75000", wa["B3"].value == 75000, str(wa["B3"].value))
check("Aannames ARPU=230", wa["B4"].value == 230, str(wa["B4"].value))
check("Aannames churn=3%", wa["B7"].value == 0.03, str(wa["B7"].value))
check("Aannames CAC=600", wa["B9"].value == 600, str(wa["B9"].value))
check("Aannames beta_einde=3", wa["B22"].value == 3, str(wa["B22"].value))
check("Aannames beta_boost=1.5", wa["B23"].value == 1.5, str(wa["B23"].value))

# Model: kernformules (MAAND=3, PFLAG=4, MKT=5, OMZET=18, RESULT=24, KAS=25, POS=31)
check("Model maand 12 = kolom N", wm.cell(row=3, column=14).value == 12, str(wm.cell(row=3, column=14).value))
check("Model beta-vlag M12", wm.cell(row=4, column=14).value == "=IF(N$3<=Aannames!$B$22,1,0)",
      str(wm.cell(row=4, column=14).value))
check("Model omzet M1", wm.cell(row=18, column=3).value == "=C15+C16+C17", str(wm.cell(row=18, column=3).value))
check("Model omzet M12", wm.cell(row=18, column=14).value == "=N15+N16+N17", str(wm.cell(row=18, column=14).value))
check("Model MRR opdrachtgevers M12 houdt rekening met beta",
      wm.cell(row=15, column=14).value == "=N10*Aannames!$B$4*(1-N4)", str(wm.cell(row=15, column=14).value))
check("Model nieuwe klanten M12 met beta-boost",
      wm.cell(row=7, column=14).value == "=ROUND(N5/Aannames!$B$9*IF(N$3<=Aannames!$B$22,Aannames!$B$23,1),0)",
      str(wm.cell(row=7, column=14).value))
check("Model resultaat M12", wm.cell(row=24, column=14).value == "=N18-N23", str(wm.cell(row=24, column=14).value))
check("Model kas M1 verwijst naar startkapitaal",
      wm.cell(row=25, column=3).value == "=Aannames!$B$3+C24", str(wm.cell(row=25, column=3).value))
check("Model kas M12 bouwt voort", wm.cell(row=25, column=14).value == "=M25+N24", str(wm.cell(row=25, column=14).value))
check("Model marketing M12 gebruikt mkt3",
      wm.cell(row=5, column=14).value == "=IF(N$3<=3,Aannames!$B$18,IF(N$3<=6,Aannames!$B$19,Aannames!$B$20))",
      str(wm.cell(row=5, column=14).value))
check("Model break-even samenvatting", "MATCH(1,C31:N31,0)" in str(wm["B35"].value), str(wm["B35"].value))
check("Model min. startkapitaal", wm["B38"].value == "=MAX(0,-MIN(C26:N26))", str(wm["B38"].value))

# Scenario: blokken en verwijzingen
check("Scenario pessimistisch CAC=1000", ws4["B4"].value == 1000, str(ws4["B4"].value))
check("Scenario base churn=3%", ws4["C3"].value == 0.03, str(ws4["C3"].value))
check("Scenario optimistisch ARPU=260", ws4["D5"].value == 260, str(ws4["D5"].value))
check("Scenario base break-even formule", ws4["B28"].value == '=IFERROR(MATCH(1,C26:N26,0),">12 maanden")', str(ws4["B28"].value))
check("Scenario base min. kapitaal", ws4["B29"].value == "=MAX(0,-MIN(C25:N25))", str(ws4["B29"].value))
check("Scenario pessimistisch min. kapitaal", ws4["B17"].value == "=MAX(0,-MIN(C13:N13))", str(ws4["B17"].value))
check("Scenario optimistisch break-even", ws4["B40"].value == '=IFERROR(MATCH(1,C38:N38,0),">12 maanden")', str(ws4["B40"].value))
check("Scenario kosten M12 base verwijst naar Model",
      str(ws4.cell(row=23, column=14).value).startswith("=Model!N$23-Model!N$21+"),
      str(ws4.cell(row=23, column=14).value))
check("Scenario omzet M12 base past beta-vlag toe",
      "(1-Model!N$4)" in str(ws4.cell(row=22, column=14).value), str(ws4.cell(row=22, column=14).value))

# Fasen: verwijzingen naar Model/Scenario
check("Fasen fase-1 prognose = Model mnd 6", wf["E4"].value == "=Model!H18", str(wf["E4"].value))
check("Fasen KPI CAC", wf["B9"].value == "=Model!N28", str(wf["B9"].value))
check("Fasen KPI LTV/CAC", wf["B10"].value == "=Model!N30", str(wf["B10"].value))
check("Fasen KPI min. kapitaal base", wf["B13"].value == "=Scenario!B29", str(wf["B13"].value))
check("Fasen KPI min. kapitaal pessimistisch", wf["B14"].value == "=Scenario!B17", str(wf["B14"].value))

# Alle formules: geldige verwijzingen (geen #REF!, geen lege formules)
bad = []
for ws_ in (wm, ws4, wf):
    for row in ws_.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                if "#REF" in c.value or c.value == "=":
                    bad.append(f"{ws_.title}!{c.coordinate}")
check("Geen kapotte formules", not bad, str(bad))

# Aannames-refs in Model/Scenario verwijzen naar bestaande rijen
missing = []
for ws_ in (wm, ws4):
    for row in ws_.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                for m in re.finditer(r"Aannames!\$B\$(\d+)", c.value):
                    r = int(m.group(1))
                    if not (3 <= r <= 23 and wa.cell(row=r, column=2).value is not None):
                        missing.append(f"{ws_.title}!{c.coordinate}->{r}")
check("Alle Aannames-refs bestaan", not missing, str(missing))

fails = [c for c in checks if not c[1]]
for name, ok, detail in checks:
    print(("PASS" if ok else "FAIL"), "-", name, ("| " + detail if detail and not ok else ""))
print(f"\n{len(checks) - len(fails)}/{len(checks)} checks geslaagd")
raise SystemExit(1 if fails else 0)
